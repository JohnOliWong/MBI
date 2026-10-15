import argparse
import os
import numpy as np
import pandas as pd
from openpyxl import Workbook


EXPORT_TO_LOG = False
EXPORT_TO_EXCEL = True

# 新指标列（需与 Train_LOSO.py 输出的 {subject}.xlsx 列名一致）
METRIC_COLS = ['B-ACC', 'precision', 'W-F1', 'kappa', 'AUC']
SELECTION_COL = 'B-ACC'   # best / top-x% 的排序依据

# 模型和数据集排序
MODEL_ORDER = ['EEGNet', 'Conformer', 'EyeFormer', 'CAFNet', 'EFNet', 'VigilanceNet', 'TSMMF', 'EFBook']
DATASET_ORDER = ['SEEDIV', 'SEEDV']

def sort_key(checkpoint_name):
    """按 dataset 优先、model 其次的固定顺序排序。"""
    model = checkpoint_name.split('_')[0]
    dataset = checkpoint_name.split('_')[1]
    return (DATASET_ORDER.index(dataset), MODEL_ORDER.index(model))

def load_subject_results(results_root, num_subjects=15):
    """读取每个受试者的 xlsx，返回 {subject: DataFrame}。"""
    results = {}
    for subject in range(1, num_subjects + 1):
        path = os.path.join(results_root, f'{subject}.xlsx')
        if os.path.exists(path):
            df = pd.read_excel(path)
            # 去掉 loss 列
            df = df.drop(columns=[c for c in df.columns if 'loss' in c.lower()], errors='ignore')
            results[subject] = df
    return results

def subject_best(df):
    """取该受试者 acc 最高的那个 epoch 的指标。"""
    idx = df[SELECTION_COL].idxmax()
    return df.loc[idx, METRIC_COLS].astype(float)

def subject_top_x(df, x):
    """取该受试者 acc 排名前 x% 的 epoch，指标求平均。"""
    n = max(1, int(np.ceil(len(df) * x)))
    top = df.nlargest(n, SELECTION_COL)
    return top[METRIC_COLS].astype(float).mean()

def subject_mean(df):
    """取该受试者所有 epoch 的指标平均。"""
    return df[METRIC_COLS].astype(float).mean()

def aggregate(results, mode, x=None):
    """对每个受试者计算一组指标，再对受试者求平均，返回一行指标。"""
    per_subject = []
    for subject, df in sorted(results.items()):
        if mode == 'best':
            per_subject.append(subject_best(df))
        elif mode == 'x%':
            per_subject.append(subject_top_x(df, x))
        elif mode == 'mean':
            per_subject.append(subject_mean(df))
    if not per_subject:
        return None
    return pd.DataFrame(per_subject)[METRIC_COLS].mean()

def format_metrics(row):
    return ', '.join(f'{col}={row[col]:.2f}' for col in METRIC_COLS)

def export_all_to_excel(excel_path, checkpoint_names, all_modes):
    """每个 checkpoint 一个列块：第一行 checkpoint_name，第二行 Mode + 指标表头，第三行起每行一个模式。"""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Source'

    n_cols = 1 + len(METRIC_COLS)
    block_height = 2 + max(
        (len([r for r in modes.values() if r is not None]) for modes in all_modes),
        default=0,
    )

    for i, (checkpoint_name, modes) in enumerate(zip(checkpoint_names, all_modes)):
        group = i // 3
        idx_in_group = i % 3
        start_col = idx_in_group * (n_cols + 1) + 1
        start_row = group * (block_height + 1) + 1

        ws.cell(row=start_row, column=start_col, value=checkpoint_name)
        ws.cell(row=start_row + 1, column=start_col, value='Mode')
        for j, col_name in enumerate(METRIC_COLS):
            ws.cell(row=start_row + 1, column=start_col + 1 + j, value=col_name)

        row_idx = start_row + 2
        for name, row in modes.items():
            if row is None:
                continue
            ws.cell(row=row_idx, column=start_col, value=name)
            for j, col_name in enumerate(METRIC_COLS):
                ws.cell(row=row_idx, column=start_col + 1 + j, value=round(float(row[col_name]), 2))
            row_idx += 1

    wb.save(excel_path)

def summarize_results(base_dir=None, checkpoint_names=None):
    """LOSO 汇总：对每个受试者 best/top-x%/mean，再对受试者求平均。返回 (汇总文本, all_modes)。

    checkpoint_names 为 None 时，扫描 Results/ 下所有 checkpoint 目录。
    """
    if base_dir is None:
        base_dir = os.path.dirname(os.path.abspath(__file__))
    results_dir = os.path.join(base_dir, 'Results')

    if checkpoint_names is None:
        checkpoint_names = [
            folder for folder in os.listdir(results_dir)
            if os.path.isdir(os.path.join(results_dir, folder)) and folder != 'Temp'
        ]
        checkpoint_names.sort(key=sort_key)

    all_modes = []
    blocks = []
    for checkpoint_name in checkpoint_names:
        dataset = checkpoint_name.split('_')[1]
        num_subjects = 15 if dataset == 'SEEDIV' else 16

        results_root = os.path.join(results_dir, checkpoint_name)
        results = load_subject_results(results_root, num_subjects)

        modes = {'best': aggregate(results, 'best')}
        for x in [0.1, 0.2, 0.3, 0.4, 0.5, 0.6, 0.7, 0.8, 0.9, 1.0]:
            modes[f'Top {x:.0%}'] = aggregate(results, 'x%', x=x)
        all_modes.append(modes)

        lines = [checkpoint_name]
        for name, row in modes.items():
            if row is None:
                continue
            lines.append(f'mode={name}')
            lines.append(f'{format_metrics(row)}')
        blocks.append('\n'.join(lines))

    text = '\n\n'.join(blocks)
    return text, all_modes


if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.abspath(__file__))
    results_dir = os.path.join(base_dir, 'Results')

    checkpoint_names = [
        folder for folder in os.listdir(results_dir)
        if os.path.isdir(os.path.join(results_dir, folder)) and folder != 'Temp'
    ]
    checkpoint_names.sort(key=sort_key)

    text, all_modes = summarize_results(base_dir=base_dir, checkpoint_names=checkpoint_names)
    print(text)

    if EXPORT_TO_LOG:
        log_path = os.path.join(base_dir, 'Results', 'Results.log')
        os.makedirs(os.path.dirname(log_path), exist_ok=True)
        with open(log_path, 'a', encoding='utf-8') as f:
            f.write(text + '\n')

    if EXPORT_TO_EXCEL:
        excel_path = os.path.join(base_dir, 'Results', 'Results_LOSO.xlsx')
        export_all_to_excel(excel_path, checkpoint_names, all_modes)
